"""Build the Word companion from the Markdown project plan using python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root/'docs/planning/PROJECT_PLAN.md'
text = source.read_text()
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(1.0)
sec.footer_distance = Inches(.35)
sec.left_margin = sec.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',28), ('Heading 1',16), ('Heading 2',13)]:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.keep_with_next=True
    style.paragraph_format.space_before=Pt(16 if name!='Title' else 0)
    style.paragraph_format.space_after=Pt(8)
# Useful page numbers for a long reference document.
footer=sec.footer.paragraphs[0]
footer.alignment=2
run=footer.add_run()
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); run._r.addnext(field)
footer.style=doc.styles['Normal']
for line in text.splitlines():
    if not line.strip(): continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('- '):
        p=doc.add_paragraph(line[2:], 'List Bullet')
        p.paragraph_format.space_after=Pt(5)
    else:
        p=doc.add_paragraph(line)
        if line.startswith('Chattanooga climbing'):
            p.runs[0].font.size=Pt(15)
        if line.startswith('Version 0.1'):
            p.runs[0].font.size=Pt(10)
            p.runs[0].font.color.rgb=RGBColor(85,85,85)
    # Preserve intentional numbered list text without introducing auto-number resets.
for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control=True
# Create real external hyperlinks for source URLs in the register.
for p in doc.paragraphs:
    if not re.search(r'https?://',p.text): continue
    raw=p.text
    p.clear()
    pos=0
    for m in re.finditer(r'https?://[^\s]+',raw):
        url=m.group().rstrip('.')
        p.add_run(raw[pos:m.start()])
        h=OxmlElement('w:hyperlink')
        rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h.set(qn('r:id'),rid)
        r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
        c=OxmlElement('w:color'); c.set(qn('w:val'),'245C65'); rp.append(c)
        r.append(rp); t=OxmlElement('w:t'); t.text=url; r.append(t); h.append(r); p._p.append(h)
        pos=m.start()+len(url)
    p.add_run(raw[pos:])
# Remove inherited decorative paragraph borders, including the default title rule.
for element in list(doc.styles.element.iter()) + list(doc.element.iter()):
    if element.tag == qn('w:pBdr'):
        element.getparent().remove(element)
doc.core_properties.title='Crag Commune project plan'
doc.core_properties.subject='Chattanooga climbing community and session planner'
doc.core_properties.author='Crag Commune contributors'
doc.core_properties.keywords='climbing, Chattanooga, community, project planning'
out=root/'docs/planning/Crag_Commune_Project_Plan.docx'
doc.save(out)
print(out)
